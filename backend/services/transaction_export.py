"""Builds the .xlsx export for the Transactions (ledger) page - the same
filtered result set the search/filter bar is showing, not just the current
page, so exporting reflects what's on screen conceptually rather than one
page of it.
"""

from __future__ import annotations

import io

from openpyxl import Workbook
from openpyxl.styles import Font

from models.transaction import Transaction

_HEADERS = [
    "Date",
    "Account",
    "Scheme",
    "Scheme Code",
    "Type",
    "Source",
    "NAV",
    "Units",
    "Amount",
    "Commission",
    "Folio",
]


def build_transactions_export(
    txns: list[Transaction], names: dict[int, tuple[str, str]]
) -> bytes:
    """names maps transaction id -> (account_name, scheme_display_name), as
    produced by services.transactions.enrich_display_names.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Transactions"

    for col, header in enumerate(_HEADERS, start=1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = Font(bold=True)

    for row_offset, txn in enumerate(txns, start=2):
        account_name, scheme_display_name = names.get(txn.id, ("—", txn.scheme_code))
        values = [
            txn.transaction_date.isoformat(),
            account_name,
            scheme_display_name,
            txn.scheme_code,
            txn.type,
            txn.source,
            float(txn.nav),
            float(txn.units),
            float(txn.amount),
            float(txn.commission),
            txn.folio_no or "",
        ]
        for col, value in enumerate(values, start=1):
            ws.cell(row=row_offset, column=col, value=value)

    for col, header in enumerate(_HEADERS, start=1):  # 11 columns, always within A-Z
        ws.column_dimensions[chr(64 + col)].width = max(12, len(header) + 4)

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
