"""End-to-end HTTP tests exercising real service wiring (not just unit
logic) - each hits the actual FastAPI routes against a throwaway SQLite DB.
"""

from datetime import date, timedelta

from models.scheme import Scheme


def _seed_scheme(db_session, code="103196", name="Nippon India ELSS Tax Saver Fund", nav="131.5970"):
    db_session.add(Scheme(scheme_code=code, scheme_name=name, latest_nav=nav, nav_date=date.today()))
    db_session.commit()


def test_health(client):
    assert client.get("/health").json() == {"status": "ok"}


def test_accounts_are_seeded(client):
    r = client.get("/api/accounts")
    assert r.status_code == 200
    names = {a["account_name"] for a in r.json()}
    assert "ICICI Direct - Ashu" in names


def test_create_purchase_then_dashboard_reflects_it(client, db_session):
    _seed_scheme(db_session)
    accounts = client.get("/api/accounts").json()
    account_id = accounts[0]["id"]

    r = client.post(
        "/api/transactions",
        json={
            "account_id": account_id,
            "scheme_code": "103196",
            "transaction_date": "2024-01-15",
            "type": "PURCHASE",
            "nav": 120.5,
            "units": 41.5,
            "amount": 5010.75,
        },
    )
    assert r.status_code == 201
    txn = r.json()
    assert txn["commission"] == 10.0  # 5010.75 - (120.5 * 41.5) = 10.0

    dash = client.get("/api/dashboard").json()
    assert dash["active_schemes"] == 1
    holding = dash["holdings"][0]
    assert holding["net_units"] == 41.5
    assert holding["net_invested"] == 5010.75  # the amount paid - not inflated by current value
    assert holding["status"] in ("In Profit", "In Loss", "Flat")


def test_net_invested_is_cost_of_units_held_not_inflated_by_current_value(client, db_session):
    """Regression test for a real bug: Net Invested was computed as
    current_value - signed_sum, which for a purchase-only holding equals
    current_value + total_purchased - roughly double the real figure (a
    $63,000 investment worth $64,582 today showed "Net Invested: $127,582").
    It must equal the actual amount paid for units still held, independent
    of today's market value.
    """
    _seed_scheme(db_session, nav="155.3240")
    account_id = client.get("/api/accounts").json()[0]["id"]

    total_invested = 0.0
    for i in range(15):
        nav = 145 + i  # a plain integer NAV
        units = 20  # a plain integer unit count -> nav * units is always exact, no rounding
        amount = nav * units
        r = client.post(
            "/api/transactions",
            json={
                "account_id": account_id,
                "scheme_code": "103196",
                "transaction_date": f"2024-{(i % 12) + 1:02d}-05",
                "type": "PURCHASE",
                "nav": nav,
                "units": units,
                "amount": amount,
            },
        )
        assert r.status_code == 201
        total_invested += amount

    dash = client.get("/api/dashboard").json()
    holding = dash["holdings"][0]

    # Current value is necessarily > 0 here (real NAV, real units) - the
    # bug only manifests when current_value is nonzero, which is every
    # normal case. Net Invested must track the ~sum of amounts paid, not
    # that plus current_value.
    assert holding["current_value"] > 0
    assert abs(holding["net_invested"] - total_invested) < 1.0
    assert holding["net_invested"] < holding["current_value"] + total_invested - 1000  # not inflated

    # % Gain must be computed against the correct (uninflated) net_invested.
    expected_pct = holding["total_gain"] / holding["net_invested"] * 100
    assert abs(holding["pct_gain"] - expected_pct) < 0.01


def test_dashboard_by_account_aggregates_at_account_level(client, db_session):
    """Two accounts, two schemes, one of the schemes held in both accounts.
    by_account must show one group per account whose schemes list is scoped
    to that account only, and whose totals (net invested, current value)
    equal the sum of that account's own scheme rows - not the portfolio-wide
    per-scheme figures repeated under both accounts.
    """
    _seed_scheme(db_session, code="103196", name="Nippon India ELSS Tax Saver Fund", nav="150.00")
    _seed_scheme(db_session, code="120503", name="Kotak Contra Fund", nav="80.00")
    accounts = client.get("/api/accounts").json()
    acc_a, acc_b = accounts[0]["id"], accounts[1]["id"]

    # Account A: 100 units of scheme 103196 @ NAV 100 = 10,000
    client.post(
        "/api/transactions",
        json={
            "account_id": acc_a,
            "scheme_code": "103196",
            "transaction_date": "2024-01-10",
            "type": "PURCHASE",
            "nav": 100,
            "units": 100,
            "amount": 10000,
        },
    )
    # Account B: 50 units of scheme 103196 @ NAV 100 = 5,000
    client.post(
        "/api/transactions",
        json={
            "account_id": acc_b,
            "scheme_code": "103196",
            "transaction_date": "2024-01-10",
            "type": "PURCHASE",
            "nav": 100,
            "units": 50,
            "amount": 5000,
        },
    )
    # Account B: 20 units of scheme 120503 @ NAV 50 = 1,000
    client.post(
        "/api/transactions",
        json={
            "account_id": acc_b,
            "scheme_code": "120503",
            "transaction_date": "2024-01-10",
            "type": "PURCHASE",
            "nav": 50,
            "units": 20,
            "amount": 1000,
        },
    )

    dash = client.get("/api/dashboard").json()
    by_account = {a["account_id"]: a for a in dash["by_account"]}

    assert acc_a in by_account and acc_b in by_account
    assert len(by_account) == 2  # accounts with no activity are omitted

    a_group = by_account[acc_a]
    assert len(a_group["schemes"]) == 1
    assert a_group["schemes"][0]["scheme_code"] == "103196"
    assert a_group["schemes"][0]["net_units"] == 100
    assert a_group["net_invested"] == 10000

    b_group = by_account[acc_b]
    assert {s["scheme_code"] for s in b_group["schemes"]} == {"103196", "120503"}
    assert b_group["net_invested"] == 6000  # 5000 + 1000

    # Each account's total is exactly the sum of its own scheme rows, not
    # the portfolio-wide per-scheme total repeated under every account
    # holding it.
    assert a_group["net_invested"] + b_group["net_invested"] == dash["net_invested"]
    assert a_group["current_value"] + b_group["current_value"] == dash["current_value"]

    # 103196's portfolio-wide net_invested (10000+5000) must not leak into
    # account A's group - it should only see its own 10000 contribution.
    portfolio_holding = next(h for h in dash["holdings"] if h["scheme_code"] == "103196")
    assert portfolio_holding["net_invested"] == 15000
    a_scheme_holding = a_group["schemes"][0]
    assert a_scheme_holding["net_invested"] == 10000


def test_net_invested_excludes_cost_of_already_redeemed_units(client, db_session):
    """After a partial redemption, Net Invested (cost of units still held)
    must drop by the FIFO cost of the units sold - not just net against the
    redemption's cash proceeds.
    """
    _seed_scheme(db_session)
    account_id = client.get("/api/accounts").json()[0]["id"]

    client.post(
        "/api/transactions",
        json={
            "account_id": account_id,
            "scheme_code": "103196",
            "transaction_date": "2024-01-01",
            "type": "PURCHASE",
            "nav": 100,
            "units": 100,
            "amount": 10000,
        },
    )
    client.post(
        "/api/transactions",
        json={
            "account_id": account_id,
            "scheme_code": "103196",
            "transaction_date": "2024-06-01",
            "type": "REDEMPTION",
            "nav": 120,
            "units": 40,
            "amount": 4800,
        },
    )

    holding = client.get("/api/dashboard").json()["holdings"][0]
    # 60 units remain at the original cost of 100/unit = 6000, regardless
    # of the redemption's own proceeds (4800) or today's NAV.
    assert holding["net_invested"] == 6000.0
    assert holding["realized_gain"] == 800.0  # 40 units * (120 - 100)


def test_dashboard_net_gain_loss_is_unrealized_and_realized_is_separate(client, db_session):
    """Net Gain/Loss is unrealized only; realized gain from redemptions is
    reported separately.
    """
    _seed_scheme(db_session)
    account_id = client.get("/api/accounts").json()[0]["id"]

    client.post(
        "/api/transactions",
        json={
            "account_id": account_id,
            "scheme_code": "103196",
            "transaction_date": "2024-01-01",
            "type": "PURCHASE",
            "nav": 100,
            "units": 100,
            "amount": 10000,
        },
    )
    client.post(
        "/api/transactions",
        json={
            "account_id": account_id,
            "scheme_code": "103196",
            "transaction_date": "2024-06-01",
            "type": "REDEMPTION",
            "nav": 120,
            "units": 40,
            "amount": 4800,
        },
    )

    dash = client.get("/api/dashboard").json()
    unrealized = sum(h["unrealized_gain"] for h in dash["holdings"])
    realized = sum(h["realized_gain"] for h in dash["holdings"])
    assert abs(dash["net_gain_loss"] - unrealized) < 0.01
    assert abs(dash["total_realized_gain"] - realized) < 0.01
    assert realized == 800.0


def test_redemption_beyond_holding_is_rejected(client, db_session):
    _seed_scheme(db_session)
    account_id = client.get("/api/accounts").json()[0]["id"]
    client.post(
        "/api/transactions",
        json={
            "account_id": account_id,
            "scheme_code": "103196",
            "transaction_date": "2024-01-15",
            "type": "PURCHASE",
            "nav": 100,
            "units": 10,
            "amount": 1000,
        },
    )
    r = client.post(
        "/api/transactions",
        json={
            "account_id": account_id,
            "scheme_code": "103196",
            "transaction_date": "2024-02-15",
            "type": "REDEMPTION",
            "nav": 110,
            "units": 20,  # more than the 10 held
            "amount": 2200,
        },
    )
    assert r.status_code == 422


def test_soft_delete_excludes_from_dashboard(client, db_session):
    _seed_scheme(db_session)
    account_id = client.get("/api/accounts").json()[0]["id"]
    txn = client.post(
        "/api/transactions",
        json={
            "account_id": account_id,
            "scheme_code": "103196",
            "transaction_date": "2024-01-15",
            "type": "PURCHASE",
            "nav": 100,
            "units": 10,
            "amount": 1000,
        },
    ).json()

    assert client.get("/api/dashboard").json()["active_schemes"] == 1
    client.delete(f"/api/transactions/{txn['id']}")
    assert client.get("/api/dashboard").json()["active_schemes"] == 0

    client.post(f"/api/transactions/{txn['id']}/restore")
    assert client.get("/api/dashboard").json()["active_schemes"] == 1


def test_sip_due_confirm_and_skip_flow(client, db_session):
    _seed_scheme(db_session)
    account_id = client.get("/api/accounts").json()[0]["id"]
    start = date.today() - timedelta(days=95)  # ~3 monthly occurrences due

    r = client.post(
        "/api/sips",
        json={
            "account_id": account_id,
            "scheme_code": "103196",
            "amount": 5000,
            "frequency": "MONTHLY",
            "start_date": start.isoformat(),
        },
    )
    assert r.status_code == 201

    due = client.get("/api/sips/due").json()
    assert len(due) >= 3

    first = due[0]
    r = client.post(
        "/api/sips/due/confirm",
        json={
            "sip_id": first["sip_id"],
            "scheduled_date": first["scheduled_date"],
            "nav": 130.0,
            "units": 38.46,
        },
    )
    assert r.status_code == 200
    assert r.json()["source"] == "SIP"

    r = client.post(
        "/api/sips/due/skip",
        json={"sip_id": due[1]["sip_id"], "scheduled_date": due[1]["scheduled_date"]},
    )
    assert r.status_code == 200

    remaining = client.get("/api/sips/due").json()
    assert len(remaining) == len(due) - 2

    # double-confirm the same occurrence is idempotent, not an error
    r2 = client.post(
        "/api/sips/due/confirm",
        json={
            "sip_id": first["sip_id"],
            "scheduled_date": first["scheduled_date"],
            "nav": 130.0,
            "units": 38.46,
        },
    )
    assert r2.status_code == 200
    assert r2.json()["id"] == r.json()["id"]  # idempotent: same transaction, not a duplicate

    # dashboard only reflects the confirmed occurrence, not the skipped one
    dash = client.get("/api/dashboard").json()
    assert dash["active_schemes"] == 1


def test_deleting_a_confirmed_sip_entry_puts_it_back_on_the_due_list(client, db_session):
    """A soft-deleted transaction must not count as 'still confirmed' -
    otherwise deleting a SIP-confirmed entry would remove it from the
    ledger/dashboard but leave it permanently missing from Entries Due too,
    with no way to re-enter it short of editing the database by hand.
    """
    _seed_scheme(db_session)
    account_id = client.get("/api/accounts").json()[0]["id"]
    start = date.today() - timedelta(days=10)

    sip = client.post(
        "/api/sips",
        json={
            "account_id": account_id,
            "scheme_code": "103196",
            "amount": 5000,
            "frequency": "MONTHLY",
            "start_date": start.isoformat(),
        },
    ).json()

    due = client.get("/api/sips/due").json()
    assert len(due) == 1
    entry = due[0]

    confirmed = client.post(
        "/api/sips/due/confirm",
        json={"sip_id": entry["sip_id"], "scheduled_date": entry["scheduled_date"], "nav": 130.0, "units": 38.46},
    ).json()
    assert client.get("/api/sips/due").json() == []  # confirmed, no longer due

    client.delete(f"/api/transactions/{confirmed['id']}")
    due_after_delete = client.get("/api/sips/due").json()
    assert len(due_after_delete) == 1  # back on the due list, not lost
    assert due_after_delete[0]["scheduled_date"] == entry["scheduled_date"]

    # re-confirming restores the same transaction rather than duplicating it
    restored = client.post(
        "/api/sips/due/confirm",
        json={"sip_id": entry["sip_id"], "scheduled_date": entry["scheduled_date"], "nav": 130.0, "units": 38.46},
    ).json()
    assert restored["id"] == confirmed["id"]
    assert restored["is_deleted"] is False
    assert client.get("/api/sips/due").json() == []
    assert sip["id"] == entry["sip_id"]


def test_commission_preview_matches_create(client):
    r = client.post(
        "/api/transactions/preview-commission",
        json={"type": "PURCHASE", "nav": 145.509, "units": 34.243, "amount": 5000},
    )
    assert r.status_code == 200
    body = r.json()
    assert abs(body["commission"] - (5000 - 145.509 * 34.243)) < 0.01


def _create_txn(client, account_id, **overrides):
    payload = {
        "account_id": account_id,
        "scheme_code": "103196",
        "transaction_date": "2024-01-15",
        "type": "PURCHASE",
        "nav": 100,
        "units": 10,
        "amount": 1000,
    }
    payload.update(overrides)
    r = client.post("/api/transactions", json=payload)
    assert r.status_code == 201, r.text
    return r.json()


def test_transaction_list_is_enriched_with_display_names(client, db_session):
    _seed_scheme(db_session)
    account_id = client.get("/api/accounts").json()[0]["id"]
    _create_txn(client, account_id)

    item = client.get("/api/transactions").json()["items"][0]
    assert item["account_name"] == "ICICI Direct - Ashu"
    assert item["scheme_display_name"] == "Nippon India ELSS Tax Saver Fund"


def test_transaction_list_pagination(client, db_session):
    _seed_scheme(db_session)
    account_id = client.get("/api/accounts").json()[0]["id"]
    for i in range(5):
        _create_txn(client, account_id, transaction_date=f"2024-01-{i + 1:02d}", amount=1000 + i, units=10)

    page1 = client.get("/api/transactions?limit=2&offset=0").json()
    assert page1["total"] == 5
    assert len(page1["items"]) == 2
    # newest first
    assert page1["items"][0]["transaction_date"] == "2024-01-05"

    page2 = client.get("/api/transactions?limit=2&offset=2").json()
    assert len(page2["items"]) == 2
    assert page2["items"][0]["transaction_date"] == "2024-01-03"

    page3 = client.get("/api/transactions?limit=2&offset=4").json()
    assert len(page3["items"]) == 1

    ids_seen = {item["id"] for page in (page1, page2, page3) for item in page["items"]}
    assert len(ids_seen) == 5  # no overlap/gaps across pages


def test_transaction_list_filters_by_account(client, db_session):
    _seed_scheme(db_session)
    accounts = client.get("/api/accounts").json()
    _create_txn(client, accounts[0]["id"])
    _create_txn(client, accounts[1]["id"])

    result = client.get(f"/api/transactions?account_id={accounts[0]['id']}").json()
    assert result["total"] == 1
    assert result["items"][0]["account_id"] == accounts[0]["id"]


def test_transaction_list_filters_by_scheme_name_free_text(client, db_session):
    _seed_scheme(db_session, code="103196", name="Nippon India ELSS Tax Saver Fund")
    _seed_scheme(db_session, code="118989", name="HDFC Mid Cap Fund")
    account_id = client.get("/api/accounts").json()[0]["id"]
    _create_txn(client, account_id, scheme_code="103196")
    _create_txn(client, account_id, scheme_code="118989")

    result = client.get("/api/transactions?scheme_name=Mid Cap").json()
    assert result["total"] == 1
    assert result["items"][0]["scheme_code"] == "118989"


def test_transaction_list_scheme_name_search_matches_users_own_display_name(client, db_session, seeded_user):
    """A user's own portfolio_schemes label (e.g. 'My Retirement Fund') must
    be searchable too, not just AMFI's own scheme_name - the search box is
    supposed to work against whatever the user actually sees on screen.
    """
    _seed_scheme(db_session, code="103196", name="Nippon India ELSS Tax Saver Fund")
    account_id = client.get("/api/accounts").json()[0]["id"]
    _create_txn(client, account_id, scheme_code="103196")

    from models.portfolio_scheme import PortfolioScheme

    db_session.add(
        PortfolioScheme(user_id=seeded_user.id, scheme_code="103196", display_name="My Retirement Fund")
    )
    db_session.commit()

    result = client.get("/api/transactions?scheme_name=Retirement").json()
    assert result["total"] == 1
    assert result["items"][0]["scheme_code"] == "103196"


def test_transaction_list_filters_by_type_and_date_range(client, db_session):
    _seed_scheme(db_session)
    account_id = client.get("/api/accounts").json()[0]["id"]
    _create_txn(client, account_id, type="PURCHASE", transaction_date="2024-01-01", units=10, amount=1000)
    _create_txn(client, account_id, type="PURCHASE", transaction_date="2024-06-01", units=5, amount=600)
    client.post(
        "/api/transactions",
        json={
            "account_id": account_id,
            "scheme_code": "103196",
            "transaction_date": "2024-07-01",
            "type": "REDEMPTION",
            "nav": 120,
            "units": 3,
            "amount": 360,
        },
    )

    redemptions = client.get("/api/transactions?type=REDEMPTION").json()
    assert redemptions["total"] == 1

    in_range = client.get("/api/transactions?date_from=2024-05-01&date_to=2024-12-31").json()
    assert in_range["total"] == 2  # the June purchase and July redemption


def test_transaction_list_excludes_deleted_unless_asked(client, db_session):
    _seed_scheme(db_session)
    account_id = client.get("/api/accounts").json()[0]["id"]
    txn = _create_txn(client, account_id)
    client.delete(f"/api/transactions/{txn['id']}")

    assert client.get("/api/transactions").json()["total"] == 0
    assert client.get("/api/transactions?include_deleted=true").json()["total"] == 1


def test_transaction_export_returns_xlsx_with_the_requested_columns(client, db_session):
    import io

    from openpyxl import load_workbook

    _seed_scheme(db_session)
    account_id = client.get("/api/accounts").json()[0]["id"]
    _create_txn(client, account_id, folio_no="F001")
    _create_txn(
        client,
        account_id,
        type="REDEMPTION",
        transaction_date="2024-02-01",
        nav=110,
        units=5,
        amount=550,
        folio_no="F001",
    )

    r = client.get("/api/transactions/export")
    assert r.status_code == 200
    assert r.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert "attachment" in r.headers["content-disposition"]

    wb = load_workbook(io.BytesIO(r.content))
    ws = wb.active
    header = [c.value for c in ws[1]]
    assert header == [
        "Date", "Account", "Scheme", "Scheme Code", "Type", "Source",
        "NAV", "Units", "Amount", "Commission", "Folio",
    ]

    rows = list(ws.iter_rows(min_row=2, values_only=True))
    assert len(rows) == 2
    # chronological ascending (oldest first), unlike the ledger's newest-first view
    assert rows[0][0] == "2024-01-15"
    assert rows[1][0] == "2024-02-01"
    assert rows[0][1] == "ICICI Direct - Ashu"
    assert rows[0][2] == "Nippon India ELSS Tax Saver Fund"
    assert rows[0][3] == "103196"
    assert rows[0][10] == "F001"  # folio


def test_transaction_export_respects_filters(client, db_session):
    import io

    from openpyxl import load_workbook

    _seed_scheme(db_session, code="103196", name="Nippon India ELSS Tax Saver Fund")
    _seed_scheme(db_session, code="118989", name="HDFC Mid Cap Fund")
    account_id = client.get("/api/accounts").json()[0]["id"]
    _create_txn(client, account_id, scheme_code="103196")
    _create_txn(client, account_id, scheme_code="118989")

    r = client.get("/api/transactions/export?scheme_name=Mid Cap")
    wb = load_workbook(io.BytesIO(r.content))
    rows = list(wb.active.iter_rows(min_row=2, values_only=True))
    assert len(rows) == 1
    assert rows[0][3] == "118989"


def test_transaction_export_includes_full_filtered_set_not_just_one_page(client, db_session):
    import io

    from openpyxl import load_workbook

    _seed_scheme(db_session)
    account_id = client.get("/api/accounts").json()[0]["id"]
    for i in range(30):  # more than one page's worth (page size is 25)
        txn_date = (date(2024, 1, 1) + timedelta(days=i)).isoformat()
        _create_txn(client, account_id, transaction_date=txn_date, units=10, amount=1000)

    r = client.get("/api/transactions/export")
    wb = load_workbook(io.BytesIO(r.content))
    rows = list(wb.active.iter_rows(min_row=2, values_only=True))
    assert len(rows) == 30


def _create_sip(client, account_id, **overrides):
    payload = {
        "account_id": account_id,
        "scheme_code": "103196",
        "amount": 5000,
        "frequency": "MONTHLY",
        "start_date": "2024-01-01",
    }
    payload.update(overrides)
    r = client.post("/api/sips", json=payload)
    assert r.status_code == 201, r.text
    return r.json()


def test_sip_list_is_enriched_with_display_names(client, db_session):
    _seed_scheme(db_session)
    account_id = client.get("/api/accounts").json()[0]["id"]
    _create_sip(client, account_id)

    sip = client.get("/api/sips").json()[0]
    assert sip["account_name"] == "ICICI Direct - Ashu"
    assert sip["scheme_display_name"] == "Nippon India ELSS Tax Saver Fund"


def test_sip_hold_stops_due_entries_and_resume_brings_them_back(client, db_session):
    _seed_scheme(db_session)
    account_id = client.get("/api/accounts").json()[0]["id"]
    start = date.today() - timedelta(days=40)
    sip = _create_sip(client, account_id, start_date=start.isoformat())

    assert len(client.get("/api/sips/due").json()) >= 1

    held = client.post(f"/api/sips/{sip['id']}/hold").json()
    assert held["is_active"] is False
    assert held["end_date"] is None  # held, not stopped - no end recorded
    assert client.get("/api/sips/due").json() == []

    # a held SIP still exists but active_only=true hides it by default
    assert client.get("/api/sips").json() == []
    assert len(client.get("/api/sips?active_only=false").json()) == 1

    resumed = client.post(f"/api/sips/{sip['id']}/resume").json()
    assert resumed["is_active"] is True
    assert resumed["id"] == sip["id"]  # same row, not a new version
    assert len(client.get("/api/sips/due").json()) >= 1


def test_sip_stop_records_end_date_and_removes_it_from_due(client, db_session):
    _seed_scheme(db_session)
    account_id = client.get("/api/accounts").json()[0]["id"]
    start = date.today() - timedelta(days=40)
    sip = _create_sip(client, account_id, start_date=start.isoformat())

    stopped = client.post(f"/api/sips/{sip['id']}/stop", json={}).json()
    assert stopped["is_active"] is False
    assert stopped["end_date"] == date.today().isoformat()
    assert client.get("/api/sips/due").json() == []
    assert client.get("/api/sips/upcoming").json() == []


def test_sip_stop_with_explicit_end_date_before_start_is_rejected(client, db_session):
    _seed_scheme(db_session)
    account_id = client.get("/api/accounts").json()[0]["id"]
    sip = _create_sip(client, account_id, start_date="2024-06-01")

    r = client.post(f"/api/sips/{sip['id']}/stop", json={"end_date": "2024-01-01"})
    assert r.status_code == 422


def test_sip_resume_clears_a_prior_stop_end_date(client, db_session):
    """Resuming a stopped SIP without clearing end_date would reactivate it
    into a state with no future occurrences at all (start_date > end_date),
    which would look identical to still being paused.
    """
    _seed_scheme(db_session)
    account_id = client.get("/api/accounts").json()[0]["id"]
    sip = _create_sip(client, account_id, start_date="2024-01-01")
    client.post(f"/api/sips/{sip['id']}/stop", json={"end_date": "2024-06-01"})

    resumed = client.post(f"/api/sips/{sip['id']}/resume").json()
    assert resumed["is_active"] is True
    assert resumed["end_date"] is None


def test_sip_edit_via_version_creates_a_new_active_row_and_enriches_it(client, db_session):
    _seed_scheme(db_session)
    account_id = client.get("/api/accounts").json()[0]["id"]
    sip = _create_sip(client, account_id, amount=5000)

    new_version = client.post(
        f"/api/sips/{sip['id']}/version",
        json={"effective_date": date.today().isoformat(), "amount": 6000},
    ).json()

    assert new_version["id"] != sip["id"]
    assert new_version["amount"] == 6000.0
    assert new_version["parent_sip_id"] == sip["id"]
    assert new_version["is_active"] is True
    assert new_version["account_name"] == "ICICI Direct - Ashu"
    assert new_version["scheme_display_name"] == "Nippon India ELSS Tax Saver Fund"

    old = next(s for s in client.get("/api/sips?active_only=false").json() if s["id"] == sip["id"])
    assert old["is_active"] is False
    assert old["end_date"] is not None


def test_display_names_include_plan_and_option_when_no_curated_override_exists(client, db_session):
    """Regression test: a manually-entered transaction (the common case -
    nothing seeds portfolio_schemes for it) must show the complete name -
    'Kotak Contra Fund - Regular Plan - Growth', not just 'Kotak Contra
    Fund' - everywhere a scheme's display name appears: the ledger, the
    dashboard breakdown, and the SIP list alike.
    """
    db_session.add(
        Scheme(
            scheme_code="998877",
            scheme_name="Kotak Contra Fund",
            plan="Regular Plan",
            option="Growth",
            latest_nav="155.3240",
            nav_date=date.today(),
        )
    )
    db_session.commit()
    account_id = client.get("/api/accounts").json()[0]["id"]
    expected = "Kotak Contra Fund - Regular Plan - Growth"

    txn = _create_txn(client, account_id, scheme_code="998877")
    assert txn["scheme_display_name"] == expected

    listed = client.get("/api/transactions").json()["items"][0]
    assert listed["scheme_display_name"] == expected

    holding = client.get("/api/dashboard").json()["holdings"][0]
    assert holding["display_name"] == expected

    sip = _create_sip(client, account_id, scheme_code="998877")
    assert sip["scheme_display_name"] == expected


def test_nav_history_endpoint_returns_null_not_404_when_not_found(client, monkeypatch):
    # Patched where it's *used* (api.routes.nav did `from services.nav_history
    # import fetch_nav_on_or_before`, binding its own name) - not where it's
    # defined, which patching services.nav_history itself would miss entirely.
    import api.routes.nav as nav_route

    monkeypatch.setattr(nav_route, "fetch_nav_on_or_before", lambda code, on: None)
    r = client.get("/api/schemes/103196/nav-history?on=2024-06-14")
    assert r.status_code == 200
    assert r.json() is None


def test_nav_history_endpoint_returns_prefill_data(client, monkeypatch):
    from decimal import Decimal

    import api.routes.nav as nav_route
    from services.nav_history import NavHistoryPoint

    monkeypatch.setattr(
        nav_route,
        "fetch_nav_on_or_before",
        lambda code, on: NavHistoryPoint(nav_date=date(2024, 6, 14), nav=Decimal("111.2233")),
    )
    r = client.get("/api/schemes/103196/nav-history?on=2024-06-15")
    assert r.status_code == 200
    body = r.json()
    assert body["nav"] == 111.2233
    assert body["nav_date"] == "2024-06-14"


def test_sip_confirm_with_transaction_date_override(client, db_session):
    """The user can override the system-computed trading-day-adjusted date
    with whatever their actual statement shows - the idempotency key
    (scheduled_date) stays fixed either way.
    """
    _seed_scheme(db_session)
    account_id = client.get("/api/accounts").json()[0]["id"]
    start = date.today() - timedelta(days=10)
    client.post(
        "/api/sips",
        json={
            "account_id": account_id,
            "scheme_code": "103196",
            "amount": 5000,
            "frequency": "MONTHLY",
            "start_date": start.isoformat(),
        },
    )
    entry = client.get("/api/sips/due").json()[0]
    override_date = (date.fromisoformat(entry["adjusted_date"]) + timedelta(days=2)).isoformat()

    r = client.post(
        "/api/sips/due/confirm",
        json={
            "sip_id": entry["sip_id"],
            "scheduled_date": entry["scheduled_date"],
            "transaction_date": override_date,
            "nav": 130.0,
            "units": 38.46,
        },
    )
    assert r.status_code == 200
    assert r.json()["transaction_date"] == override_date

    # re-confirming the same occurrence (same scheduled_date) is still idempotent
    r2 = client.post(
        "/api/sips/due/confirm",
        json={
            "sip_id": entry["sip_id"],
            "scheduled_date": entry["scheduled_date"],
            "transaction_date": override_date,
            "nav": 130.0,
            "units": 38.46,
        },
    )
    assert r2.json()["id"] == r.json()["id"]


def test_account_update(client):
    account = client.get("/api/accounts").json()[0]
    r = client.patch(
        f"/api/accounts/{account['id']}", json={"account_name": "Renamed Account", "platform": "New Platform"}
    )
    assert r.status_code == 200
    body = r.json()
    assert body["id"] == account["id"]
    assert body["account_name"] == "Renamed Account"
    assert body["platform"] == "New Platform"

    listed = client.get("/api/accounts").json()
    assert any(a["account_name"] == "Renamed Account" for a in listed)


def test_account_create_and_update_reject_duplicate_names(client):
    accounts = client.get("/api/accounts").json()
    existing_name = accounts[0]["account_name"]

    r = client.post("/api/accounts", json={"account_name": existing_name, "platform": "X"})
    assert r.status_code == 422

    r2 = client.patch(f"/api/accounts/{accounts[1]['id']}", json={"account_name": existing_name, "platform": "X"})
    assert r2.status_code == 422


def test_account_update_not_found(client):
    r = client.patch("/api/accounts/999999", json={"account_name": "X", "platform": "Y"})
    assert r.status_code == 404


def test_nav_sync_endpoint_includes_history_result_fields(client, monkeypatch):
    # Patched where used: api.routes.nav imports both sync_navs and
    # sync_nav_history by name into its own module namespace.
    import api.routes.nav as nav_route
    from services.amfi import SyncResult
    from services.portfolio_history import HistorySyncResult

    monkeypatch.setattr(
        nav_route,
        "sync_navs",
        lambda db: SyncResult(rows_parsed=5, rows_upserted=5, rows_skipped=0, feed_date=date(2024, 6, 14), duration_ms=10),
    )
    monkeypatch.setattr(
        nav_route,
        "sync_nav_history",
        lambda db, user: HistorySyncResult(schemes_synced=2, schemes_failed=1, points_inserted=400, duration_ms=250),
    )
    r = client.post("/api/nav/sync")
    assert r.status_code == 200
    body = r.json()
    assert body["rows_upserted"] == 5
    assert body["history_schemes_synced"] == 2
    assert body["history_schemes_failed"] == 1
    assert body["history_points_inserted"] == 400
    assert body["history_duration_ms"] == 250


def test_dashboard_history_empty_before_any_transactions(client):
    r = client.get("/api/dashboard/history")
    assert r.status_code == 200
    assert r.json() == []


def test_dashboard_history_endpoint_reflects_transactions_and_synced_nav(client, db_session):
    from models.scheme_nav_history import SchemeNavHistory

    _seed_scheme(db_session, code="103196", nav="150.0000")
    account_id = client.get("/api/accounts").json()[0]["id"]

    r = client.post(
        "/api/transactions",
        json={
            "account_id": account_id,
            "scheme_code": "103196",
            "transaction_date": "2024-01-10",
            "type": "PURCHASE",
            "nav": 100,
            "units": 10,
            "amount": 1000,
        },
    )
    assert r.status_code == 201

    db_session.add(SchemeNavHistory(scheme_code="103196", nav_date=date(2024, 1, 10), nav="100.0000"))
    db_session.add(SchemeNavHistory(scheme_code="103196", nav_date=date(2024, 1, 12), nav="120.0000"))
    db_session.commit()

    r = client.get("/api/dashboard/history")
    assert r.status_code == 200
    series = r.json()
    by_date = {p["as_of"]: p for p in series}

    assert series[0]["as_of"] == "2024-01-10"
    assert by_date["2024-01-10"]["current_value"] == 1000.0  # 10 units * NAV 100
    assert by_date["2024-01-10"]["net_invested"] == 1000.0  # cost basis - what was actually paid
    assert by_date["2024-01-10"]["pct_gain"] == 0.0  # bought at cost, NAV unchanged on day 1

    assert by_date["2024-01-11"]["current_value"] == 1000.0  # forward-filled (no NAV published that day)

    assert by_date["2024-01-12"]["current_value"] == 1200.0  # 10 units * NAV 120
    assert by_date["2024-01-12"]["net_invested"] == 1000.0  # cost basis unchanged - no new purchase
    assert by_date["2024-01-12"]["pct_gain"] == 20.0  # (1200 - 1000) / 1000 * 100

    # series runs all the way to today
    assert series[-1]["as_of"] == date.today().isoformat()
